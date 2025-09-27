# gerar_dossie.py
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# Criar documento
doc = Document()

# Função para adicionar título
def add_title(text, size=16, bold=True, center=True):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if center:
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    return p

# Função para adicionar parágrafo
def add_paragraph(text, size=12, bold=False, center=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if center:
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    return p

# -----------------------------
# Capa
# -----------------------------
add_title("Dossiê Smart Office", size=20)
add_paragraph("Aluno: Beatriz", size=14)
add_paragraph("Disciplina: Gestão de Projetos", size=14)
add_paragraph("Data: 27/09/2025", size=14)
doc.add_page_break()

# -----------------------------
# Fase 1 – Análise de Riscos
# -----------------------------
add_title("Fase 1 – Análise de Riscos e RACI", size=16, center=False)
doc.add_paragraph("1. Tabela de Riscos:")

# Riscos
riscos = [
    ["Falha de integração dos sensores com a rede", "Técnicos", "Testes de integração em laboratório; escolher sensores com SDKs bem documentados", "Média", "Alto"],
    ["Latência/perda de pacotes na rede Wi-Fi", "Técnicos", "VLAN dedicada para IoT, QoS, testes de cobertura e repetidores", "Média", "Médio"],
    ["Dados corruptos ou formatos inconsistentes", "Técnicos", "Definir esquema (JSON/CSV) e validar na ingestão; pipeline de ETL", "Baixa", "Médio"],
    ["Atraso na entrega de hardware (fornecedor)", "Operacional", "Planejar buffer de prazo; múltiplos fornecedores; contrato com SLA", "Média", "Médio"],
    ["Falta de adesão dos funcionários (resistência)", "Humanos", "Comunicação clara, demonstrações e piloto; feedback loop com RH", "Média", "Médio"],
    ["Questões de privacidade / conformidade (dados pessoais)", "Humanos", "Anonimizar dados, revisar LGPD, definir DPO e políticas de retenção", "Baixa", "Alto"],
    ["Consumo de energia inesperado pelos sensores", "Operacional", "Escolher sensores de baixo consumo; monitorar consumo; configurar sleep", "Baixa", "Médio"],
    ["Manutenção física (danos/roubo dos sensores)", "Operacional", "Localização segura, travas, inventário e seguros", "Baixa", "Médio"],
    ["Falha no algoritmo de detecção de ocupação (falsos positivos)", "Técnicos", "Validar com dados reais, calibrar thresholds, logs para retraining", "Média", "Médio"],
    ["Excesso de dependência de um fornecedor único", "Operacional", "Estratégia multi-vendor, cláusulas contratuais", "Média", "Alto"]
]

table = doc.add_table(rows=1, cols=5)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "Risco"
hdr_cells[1].text = "Categoria"
hdr_cells[2].text = "Mitigação"
hdr_cells[3].text = "Probabilidade"
hdr_cells[4].text = "Impacto"

for risco in riscos:
    row_cells = table.add_row().cells
    for i in range(5):
        row_cells[i].text = risco[i]

doc.add_paragraph("\n2. Matriz RACI:")

# Matriz RACI
atividades = ["Seleção de Fornecedores", "Instalação Física", "Desenvolvimento do Dashboard",
              "Comunicação com Funcionários", "Configuração Rede & Segurança", "Execução do Piloto"]
stakeholders = ["Diretoria TI", "Facilities", "RH", "Diretoria Financeira", "Funcionários", "Gerente Projeto"]
raci = [
    ["C","R","C","A","I","R"],
    ["C","R","I","I","I","A"],
    ["R","I","C","C","I","A"],
    ["C","I","R","I","I","A"],
    ["A","C","I","I","I","R"],
    ["A","R","C","I","C","R"]
]

table = doc.add_table(rows=1, cols=7)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "Atividade / Stakeholder"
for i in range(6):
    hdr_cells[i+1].text = stakeholders[i]

for idx, act in enumerate(atividades):
    row_cells = table.add_row().cells
    row_cells[0].text = act
    for j in range(6):
        row_cells[j+1].text = raci[idx][j]

doc.add_page_break()

# -----------------------------
# Fase 2 – Plano de Comunicação
# -----------------------------
add_title("Fase 2 – Plano de Comunicação", size=16, center=False)
comunicacoes = [
    ["Relatório Semanal de Status","Diretoria Financeira, Diretoria TI","Semanal (toda sexta)","E-mail + PDF anexado","Gerente de Projeto"],
    ["Alerta de Manutenção","Facilities, Diretoria TI","Quando for detectado erro/queda","Slack / E-mail","Facilities"],
    ["Newsletter do Projeto","Funcionários","Mensal","Intranet + E-mail","RH / Gerente Projeto"],
    ["Convite para Piloto / Treinamento","Equipe piloto, Usuários afetados","2 semanas antes do piloto","E-mail + Reunião presencial","RH"],
    ["Relatório de ROI (Preliminar)","Diretoria Financeira","Ao término do piloto (2 semanas)","Apresentação (PPT)","Gerente Projeto / Financeiro"]
]

table = doc.add_table(rows=1, cols=5)
hdr_cells = table.rows[0].cells
headers = ["O Quê","Para Quem","Quando","Como","Responsável"]
for i in range(5):
    hdr_cells[i].text = headers[i]

for comm in comunicacoes:
    row_cells = table.add_row().cells
    for i in range(5):
        row_cells[i].text = comm[i]

doc.add_page_break()

# -----------------------------
# Fase 4 – Relatórios Automatizados
# -----------------------------
add_title("Fase 4 – Relatórios Automatizados", size=16, center=False)

add_paragraph("1. Relatório de Status:")
relatorio_status = (
    "Durante o período de 15/09/2025 a 21/09/2025, os sensores do Smart Office registraram "
    "temperaturas médias estáveis, variando entre 21,9°C e 22,5°C. "
    "A luminosidade média das salas ficou entre 412 e 418 lux. "
    "O pico de ocupação ocorreu nos escritórios, atingindo até 29 pessoas simultaneamente. "
    "Esses dados indicam eficiência no uso de energia e andamento positivo nos testes do projeto."
)
add_paragraph(relatorio_status)

add_paragraph("\n2. Ata de Reunião:")
ata_reuniao = (
    "Ata de Reunião – Projeto Smart Office\n"
    "Data: 21/09/2025\n"
    "Participantes: Equipe do projeto Smart Office\n\n"
    "Assuntos discutidos:\n"
    "1. Resultados dos sensores simulados: Temperatura média estável, luminosidade adequada e pico de ocupação registrado nos escritórios.\n"
    "2. Observações: picos inesperados de ocupação em determinados dias.\n"
    "3. Decisões tomadas:\n"
    "- Investigar o script de simulação do sensor de ocupação (responsável: Beatriz)\n"
    "- Preparar apresentação dos resultados iniciais para o RH (responsável: Beatriz)"
)
add_paragraph(ata_reuniao)

# -----------------------------
# Salvar documento
# -----------------------------
doc.save("Dossie_SmartOffice_Beatriz.docx")
print("Arquivo gerado: Dossie_SmartOffice_Beatriz.docx")
